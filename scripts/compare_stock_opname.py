#!/usr/bin/env python3
"""
compare_stock_opname.py

Compares cleaned ERP Stock Quant data (the output of clean_stock_quant.py) for one
city against a physical Stock Opname count for the same city/date, and produces a
5-sheet Excel report: Summary, Product Summary, IMEI Mismatch, Aging Report,
Aging Report Detail.

## Inputs
- --cleaned-csv     : output of clean_stock_quant.py (already has an Abnormality column)
- --opname-xlsx     : the physical count workbook. Must contain a
                       "Stock <CITY> <DATE> - IMEI" sheet and a
                       "Stock <CITY> <DATE> - Non IMEI" sheet for the resolved date.
- --exclusion-xlsx  : the exclusion workbook (single sheet, single "Product"
                       column of names to drop from scope -- e.g. "Stock
                       Opname Exclude Item.xlsx"). Also requires ODOO_URL /
                       ODOO_DB / ODOO_USERNAME / ODOO_API_KEY to be set (see
                       odoo_client.py) -- scope and price are pulled live from
                       Odoo, not from a local Inventory Masterfile anymore.
- --city            : e.g. "SBY", "JKT"
- --date            : optional 6-digit override (e.g. "260730"); if omitted it's
                       inferred from the --cleaned-csv filename.

## Product scope and pricing (replaces the Inventory Masterfile, 2026-09-22)

The Inventory Masterfile used to BE the whitelist (only products listed in
its "Category" sheet were in scope) and the source of Price/Category/New
Group. That's now split:
  - Whitelist: every Odoo `product.template` in `odoo_client.CATEGORY_IDS`
    (currently `[34]`), minus anything named in --exclusion-xlsx. See
    `odoo_client.fetch_in_scope_products`.
  - Price: the minimum `product.supplierinfo` price per product, converted to
    IDR via `odoo_client.CURRENCY_TO_IDR_RATE`. See
    `odoo_client.fetch_product_info`.
  - Category / New Group - July 25: both now get the SAME price-based tier
    (A/B/C) from `odoo_client.category_from_price` -- C <=150,000,
    B 150,001-600,000, A >=600,001 IDR. Placeholder bands given 2026-09-22 by
    the business owner, explicitly confirmed changeable. This single tier
    feeds both columns (see `load_scope_and_prices`), which is what makes
    "Stock Opname Accuracy (A+B)", "Discrepancy Value (A+B)", and the
    "Summary per Group (A-C)" section work again.

## Key design decisions (locked in with the business owner on 2026-08-11)

1. ERP "on-hand" rows that carry ANY (non-empty) Abnormality flag from
   clean_stock_quant.py are EXCLUDED from BOTH the IMEI-lot match and the
   non-IMEI qty sum. This was `ERP Duplicate line` only through 2.0.0; broadened
   to every flag type on 2026-08-26 after SBY 260826 data showed the narrower
   rule let other abnormalities distort results: a `Negative Qty` + `Owner
   Mismatch` line (qty -40, owner tagged to a courier company, not the
   warehouse) inflated the non-IMEI qty comparison for ULTRASONIC FUEL SENSOR
   THINKSONIC TUB01 into a -39 unit / -41.3M IDR "discrepancy" that vanished
   once that line and one other Owner-Mismatch line were excluded — the
   remaining clean ERP qty (3) matched the physical count (3) exactly. Three
   ERP-only IMEIs in that same run were likewise all already flagged
   (Owner Mismatch and/or Imei Length Difference/Alphanumeric IMEI). Original
   `ERP Duplicate line` rationale still holds as a subset: these are lines
   where the same lot has Quantity > 0 in two places at once (e.g. still shown
   in `SBY/Stock` *and* already in `Partners/Customers`) — a physical count
   independently confirms these are phantom ERP bookings.
2. IMEI matching is done by **Lot/Serial Number alone** (not Product+Lot), so a unit
   that ended up recorded under a different product name on one side still counts as
   a match — but gets flagged as "Product Mismatch" rather than silently ignored.
3. Stock Opname Accuracy has two independent measures, both computed for "All
   products" and "A+B only" (New Group A/B; Group C excluded from the A+B cut):
     a. Qty-based:      1 - (sum|Qty_ERP - Qty_Physical| / sum(Qty_ERP))
     b. Product-count-based, IMEI-tracked products only: how many IMEI-tracked
        products have ALL their IMEIs matched ("cocok") vs at least one mismatch
        ("tidak cocok").
4. Product-name matching against the whitelist is CASE-INSENSITIVE
   (`build_product_lookup`) — confirmed against real SBY 260826 data where the
   ERP export had `FUSE 2A` (all caps) against the masterfile's `Fuse 2A`
   (this predates the Odoo-based whitelist but the same casing drift applies
   to Odoo product names too). Case-sensitive matching silently dropped that
   row out of scope entirely (treated as ERP qty 0 rather than compared),
   producing a false discrepancy even though the row itself had no
   Abnormality flag and its qty actually matched the physical count. Every
   ERP/physical row is normalized to the whitelist's canonical casing before
   the whitelist check.
5. Aging Report / Aging Report Detail: age is the whole number of months from
   each ERP row's "Lot Created Date" to the opname date (day-of-month aware:
   a lot created on the 20th isn't yet "1 month old" on the 15th of next
   month), bucketed into < 3 Mo / 3-6 Mo / 6-9 Mo / 9-12 Mo / > 12 Mo. Built
   from the same scoped, non-abnormal ERP rows used everywhere else in this
   report (city's */Stock location, whitelisted product, no Abnormality flag)
   — so a row already excluded from the Product Summary/IMEI Mismatch sheets
   is excluded here too. Rows with an unparseable/blank Lot Created Date are
   skipped from "Aging Report" (can't bucket them) but still listed in "Aging
   Report Detail" with a blank Age/Aging so they're visible rather than
   silently dropped. Ported from a business-owner-provided reference notebook
   ("v2_stock_opname.ipynb") where this logic existed but was never wired up.
"""

import argparse
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

import csv
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

import odoo_client

# Some log/report strings below use an em dash (—). On Windows, stdout defaults
# to the console's codepage (e.g. cp1252) rather than UTF-8, which raises
# UnicodeEncodeError on that character — force UTF-8 explicitly.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--cleaned-csv", required=True, help="Output of clean_stock_quant.py")
    p.add_argument("--opname-xlsx", required=True, help="Physical Stock Opname workbook")
    p.add_argument("--exclusion-xlsx", required=True, help="Exclusion workbook (single 'Product' column of names to drop from the Odoo category-34 scope)")
    p.add_argument("--city", required=True, help="e.g. SBY, JKT")
    p.add_argument("--date", default=None, help="6-digit date token e.g. 260730; inferred from --cleaned-csv filename if omitted")
    p.add_argument("--output", default=None)
    return p.parse_args()


def infer_date(cleaned_csv_path, override):
    if override:
        return override
    m = re.search(r"(\d{6})", Path(cleaned_csv_path).stem)
    if not m:
        raise ValueError(f"Could not infer a 6-digit date from cleaned CSV filename: {cleaned_csv_path}. Pass --date explicitly.")
    return m.group(1)


def load_scope_and_prices(exclusion_xlsx_path):
    """Replaces load_masterfile(): whitelist + price now come live from Odoo
    (see the "Product scope and pricing" section of this module's docstring),
    not from a local Inventory Masterfile.

    Whitelist keys are odoo_client.canonical_name() -- "[code] Name" for
    products with a default_code, else plain name -- which is the form the
    ERP/physical Stock Opname data actually uses for bracket-coded devices.
    Looking up price by product.template id (not by re-searching on name)
    avoids re-deriving that bracket form a second time.

    category and new_group both get the SAME value: odoo_client's price-tier
    (A/B/C, see category_from_price -- placeholder bands, confirmed
    changeable). There's no other classification source right now, and the
    tier's own A/B/C labels are exactly what "New Group - July 25" already
    meant (it's what drives Accuracy (A+B), Discrepancy (A+B), and the
    Summary per Group sheet section) -- so reusing it for both avoids a
    product showing a "Category" that contradicts its own group."""
    records = odoo_client.fetch_in_scope_products(exclusion_xlsx_path)
    price_by_tmpl_id = odoo_client.fetch_price_by_tmpl_ids([r["id"] for r in records])
    master = {}
    for r in records:
        # A product with no usable supplierinfo price (price_by_tmpl_id has no
        # entry for it) is treated as price=0 for BOTH the report's Price
        # column and its tier -- so it lands in Category C rather than going
        # unclassified (None) in the "Category"/"New Group" columns.
        price = price_by_tmpl_id.get(r["id"], 0)
        tier = odoo_client.category_from_price(price)
        master[r["canonical_name"]] = {"price": price, "category": tier, "new_group": tier}
    return master


def build_product_lookup(whitelist):
    """Case-insensitive product-name lookup: lowercased name -> canonical
    (Odoo) casing. ERP exports and the physical Opname workbook have both
    been seen to drift in casing from the canonical product name (e.g.
    'FUSE 2A' in the ERP vs 'Fuse 2A' in Odoo) — matched case-sensitively,
    that silently drops the row out of scope entirely rather than comparing
    it, so every row is normalized to the whitelist's casing before the
    whitelist check runs."""
    return {p.strip().lower(): p for p in whitelist}


def load_cleaned_erp(path, city, product_lookup):
    """Returns:
      erp_lot_rows: lot -> ERP row dict (qty>0, NOT flagged with any Abnormality)
      qty_by_product: product -> sum(Quantity) across all scoped, non-flagged rows
                       (used for non-IMEI products)
      abnormality_excluded_count: how many rows were excluded due to having any
                       (non-empty) Abnormality flag
      lot_tracked_products: products that have >=1 lot-bearing row, regardless
                       of exclusion
      clean_rows: every scoped, non-flagged row (kept whole, incl. "Lot
                       Created Date") — feeds the Aging Report/Detail sheets
    """
    location = f"{city}/Stock"
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    scoped = []
    for r in rows:
        if r["Location"] != location:
            continue
        canonical = product_lookup.get(r["Product"].strip().lower())
        if canonical is None:
            continue
        r["Product"] = canonical
        scoped.append(r)

    erp_lot_rows = {}
    qty_by_product = defaultdict(float)
    lot_tracked_products = set()  # products that have >=1 lot-bearing row, REGARDLESS of exclusion
    abnormality_excluded_count = 0
    clean_rows = []
    for r in scoped:
        lot = r["Lot/Serial Number"].strip()
        if lot:
            lot_tracked_products.add(r["Product"])
        if r["Abnormality"]:
            abnormality_excluded_count += 1
            continue
        try:
            qty = float(r["Quantity"])
        except ValueError:
            qty = 0.0
        qty_by_product[r["Product"]] += qty
        if lot and qty > 0:
            erp_lot_rows[lot] = r
        clean_rows.append(r)

    return erp_lot_rows, qty_by_product, abnormality_excluded_count, lot_tracked_products, clean_rows


def load_physical_imei(wb, sheet_name, product_lookup):
    ws = wb[sheet_name]
    physical_lot_rows = {}
    scan_counts = Counter()
    for row in ws.iter_rows(min_row=2, values_only=True):
        product, sku, lokasi, box_ke, lot, nomor_box = (tuple(row) + (None,) * 6)[:6]
        if product is None:
            continue
        product = product_lookup.get(str(product).strip().lower())
        if product is None:
            continue
        if lot is None or str(lot).strip() == "":
            continue
        lot = str(lot).strip()
        scan_counts[lot] += 1
        physical_lot_rows[lot] = {
            "Product": product, "SKU": sku, "Lokasi Gudang": lokasi,
            "Box Ke-": box_ke, "Nomor Box": nomor_box,
        }
    duplicate_scans = {lot: n for lot, n in scan_counts.items() if n > 1}
    return physical_lot_rows, duplicate_scans


def load_physical_non_imei(wb, sheet_name, product_lookup):
    ws = wb[sheet_name]
    qty_by_product = defaultdict(float)
    for row in ws.iter_rows(min_row=2, values_only=True):
        product, sku, lokasi, box_ke, qty, nomor_box = (tuple(row) + (None,) * 6)[:6]
        if product is None:
            continue
        product = product_lookup.get(str(product).strip().lower())
        if product is None:
            continue
        try:
            qty_by_product[product] += float(qty or 0)
        except (TypeError, ValueError):
            pass
    return qty_by_product


def build_product_summary(whitelist, master, erp_lot_rows, erp_qty_by_product,
                           physical_lot_rows, physical_qty_nonimei, matched_lots, lot_tracked_products):
    """One row per whitelisted product that has any ERP or physical activity.

    A product counts as "IMEI-tracked" if it EVER shows up with a lot number on
    either side — even if every single one of its ERP lots got excluded as
    ERP Duplicate line (that should net out to a 0-vs-0 match, not fall back to
    the raw/uncorrected quantity sum, which would silently re-introduce the very
    phantom units the exclusion was meant to remove)."""
    imei_products = lot_tracked_products | {r["Product"] for r in physical_lot_rows.values()}

    rows = []
    for product in sorted(whitelist):
        info = master[product]
        if product in imei_products:
            qty_erp = sum(1 for r in erp_lot_rows.values() if r["Product"] == product)
            qty_phys = sum(1 for r in physical_lot_rows.values() if r["Product"] == product)
            relevant_lots = {lot for lot, r in erp_lot_rows.items() if r["Product"] == product} | \
                            {lot for lot, r in physical_lot_rows.items() if r["Product"] == product}
            matched_for_product = len(relevant_lots & matched_lots)
            match_rate = (matched_for_product / len(relevant_lots)) if relevant_lots else None
        else:
            qty_erp = erp_qty_by_product.get(product, 0)
            qty_phys = physical_qty_nonimei.get(product, 0)
            match_rate = None

        if qty_erp == 0 and qty_phys == 0:
            continue  # nothing to report for a product with zero activity on both sides

        difference = qty_erp - qty_phys
        rows.append({
            "Product": product,
            "Price in IDR": info["price"],
            "Qty_ERP": qty_erp,
            "Qty_Physical": qty_phys,
            "Difference": difference,
            "Discrepancy Value (IDR)": difference * (info["price"] or 0),
            "IMEI Match Rate (%)": match_rate,
            "Category": info["category"],
            "New Group - July 25": info["new_group"],
        })
    return rows


def compute_accuracy(product_rows, groups=None):
    rows = [r for r in product_rows if groups is None or r["New Group - July 25"] in groups]
    total_qty_erp = sum(r["Qty_ERP"] for r in rows)
    total_abs_diff = sum(abs(r["Difference"]) for r in rows)
    if total_qty_erp == 0:
        return None
    return 1 - (total_abs_diff / total_qty_erp)


# Whole-month age buckets. Each tuple is (inclusive upper bound in months, label);
# the last bucket catches everything above the previous bound.
AGING_BUCKETS = [
    (2, "< 3 Mo"),
    (5, "3-6 Mo"),
    (8, "6-9 Mo"),
    (11, "9-12 Mo"),
    (float("inf"), "> 12 Mo"),
]
AGING_ORDER = {label: i for i, (_, label) in enumerate(AGING_BUCKETS)}
AGING_COLORS = {
    "< 3 Mo": ("C6EFCE", "276221"),
    "3-6 Mo": ("FFEB9C", "7D6608"),
    "6-9 Mo": ("FFCC7A", "7D4A00"),
    "9-12 Mo": ("FFC7CE", "9C0006"),
    "> 12 Mo": ("FF0000", "FFFFFF"),
}
LOT_DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y")


def parse_lot_created_date(raw):
    raw = (raw or "").strip()
    if not raw:
        return None
    for fmt in LOT_DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def age_in_months(created, as_of):
    """Whole months between `created` and `as_of`, day-of-month aware: a lot
    created on the 20th isn't yet a full month old on the 15th of next month."""
    months = (as_of.year - created.year) * 12 + (as_of.month - created.month)
    if created.day > as_of.day:
        months -= 1
    return max(months, 0)


def age_bucket(months):
    for upper, label in AGING_BUCKETS:
        if months <= upper:
            return label
    return AGING_BUCKETS[-1][1]


def build_aging_report(clean_rows, master, as_of_date):
    """Product x Aging-bucket rows, summed qty. Rows with no parseable Lot
    Created Date are skipped (can't be bucketed) and counted in `skipped`."""
    totals = defaultdict(float)
    skipped = 0
    for r in clean_rows:
        created = parse_lot_created_date(r.get("Lot Created Date"))
        if created is None:
            skipped += 1
            continue
        try:
            qty = float(r["Quantity"])
        except ValueError:
            qty = 0.0
        totals[(r["Product"], age_bucket(age_in_months(created, as_of_date)))] += qty

    rows = [
        {
            "Product": product,
            "Category": master.get(product, {}).get("category"),
            "Qty_ERP": qty,
            "Aging": bucket,
        }
        for (product, bucket), qty in totals.items()
    ]
    rows.sort(key=lambda x: (x["Product"], AGING_ORDER[x["Aging"]]))
    return rows, skipped


def build_aging_detail(clean_rows, master, as_of_date):
    """Lot-level detail: one row per (Product, Lot/Serial Number), deduplicated
    (Odoo can split one lot across multiple quant lines, e.g. partial moves) —
    summing Quantity but taking the first Lot Created Date, since it comes from
    stock.lot and is the same for every quant sharing that lot."""
    agg = {}
    for r in clean_rows:
        lot = r.get("Lot/Serial Number", "").strip()
        if not lot:
            continue
        key = (r["Product"], lot)
        try:
            qty = float(r["Quantity"])
        except ValueError:
            qty = 0.0
        if key not in agg:
            agg[key] = {"qty": 0.0, "created_raw": r.get("Lot Created Date")}
        agg[key]["qty"] += qty

    rows = []
    for (product, lot), data in agg.items():
        created = parse_lot_created_date(data["created_raw"])
        months = age_in_months(created, as_of_date) if created else None
        rows.append({
            "Product": product,
            "Lot/Serial Number": lot,
            "Category": master.get(product, {}).get("category"),
            "Lot Created Date": created,
            "Age (Months)": months,
            "Aging": age_bucket(months) if months is not None else None,
        })
    rows.sort(key=lambda x: (x["Product"], AGING_ORDER.get(x["Aging"], len(AGING_BUCKETS)), x["Lot/Serial Number"]))
    return rows


def write_report(output_path, city, date, summary_ctx, product_rows, mismatch_rows, aging_rows, aging_detail_rows):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    # ---- Summary sheet ----
    ws = wb.active
    ws.title = "Summary"
    r = 1
    ws.cell(r, 1, f"Stock Opname Summary — {city} {date}").font = Font(bold=True, size=14)
    r += 3
    ws.cell(r, 1, "IMEI Overview").font = bold
    r += 1
    for label, key in [
        ("Total IMEIs in ERP", "total_erp"),
        ("Total IMEIs in Physical", "total_physical"),
        ("Matched IMEIs", "matched"),
        ("Unmatched IMEIs in ERP", "erp_only"),
        ("Unmatched IMEIs in Physical", "physical_only"),
        ("IMEI Match Rate (%)", "match_rate"),
    ]:
        ws.cell(r, 1, label)
        ws.cell(r, 2, summary_ctx[key])
        r += 1
    r += 2
    ws.cell(r, 1, "Stock Opname Accuracy").font = bold
    r += 1
    ws.cell(r, 1, "Stock Opname Accuracy (All)"); ws.cell(r, 2, summary_ctx["accuracy_all"]); r += 1
    ws.cell(r, 1, "Stock Opname Accuracy (A+B)"); ws.cell(r, 2, summary_ctx["accuracy_ab"]); r += 1
    r += 2
    ws.cell(r, 1, "IMEI-tracked Products — Match Count").font = bold
    r += 1
    ws.cell(r, 1, "Products fully matched (cocok)"); ws.cell(r, 2, summary_ctx["products_match"]); r += 1
    ws.cell(r, 1, "Products with >=1 mismatch (tidak cocok)"); ws.cell(r, 2, summary_ctx["products_mismatch"]); r += 1
    r += 2
    ws.cell(r, 1, "Discrepancy Overview (IDR)").font = bold
    r += 1
    ws.cell(r, 1, "Discrepancy Value (All)"); ws.cell(r, 2, summary_ctx["discrepancy_all"]); r += 1
    ws.cell(r, 1, "Discrepancy Value (A+B)"); ws.cell(r, 2, summary_ctx["discrepancy_ab"]); r += 1
    r += 2
    ws.cell(r, 1, "Summary per Group (A–C)").font = bold
    r += 1
    for i, h in enumerate(["New Group - July 25", "Products_Count", "Total_Qty", "Total_Discrepancy_Qty", "Total_Rupiah"]):
        ws.cell(r, i + 1, h).font = bold
    r += 1
    for group in ["A", "B", "C"]:
        rows_g = [p for p in product_rows if p["New Group - July 25"] == group]
        ws.cell(r, 1, group)
        ws.cell(r, 2, len(rows_g))
        ws.cell(r, 3, sum(p["Qty_ERP"] for p in rows_g))
        ws.cell(r, 4, sum(abs(p["Difference"]) for p in rows_g))
        ws.cell(r, 5, sum(p["Discrepancy Value (IDR)"] for p in rows_g))
        r += 1
    for col, width in zip("ABCDE", [32, 16, 14, 20, 16]):
        ws.column_dimensions[col].width = width

    # ---- Product Summary sheet ----
    ws2 = wb.create_sheet("Product Summary")
    headers = ["Product", "Price in IDR", "Qty_ERP", "Qty_Physical", "Difference",
               "Discrepancy Value (IDR)", "IMEI Match Rate (%)", "Category", "New Group - July 25"]
    for i, h in enumerate(headers):
        ws2.cell(1, i + 1, h).font = bold
    for i, p in enumerate(product_rows):
        for j, h in enumerate(headers):
            ws2.cell(i + 2, j + 1, p[h])
    ws2.column_dimensions["A"].width = 40

    # ---- IMEI Mismatch sheet ----
    ws3 = wb.create_sheet("IMEI Mismatch")
    mismatch_headers = [
        "Lot/Serial Number", "Match Status",
        "Product_ERP", "Location", "Quantity", "Owner", "Lot Created Date", "Abnormality",
        "Product_Physical", "SKU", "Lokasi Gudang", "Box Ke-", "Nomor Box",
    ]
    for i, h in enumerate(mismatch_headers):
        ws3.cell(1, i + 1, h).font = bold
    for i, row in enumerate(mismatch_rows):
        for j, h in enumerate(mismatch_headers):
            ws3.cell(i + 2, j + 1, row.get(h))
    ws3.column_dimensions["A"].width = 20
    ws3.column_dimensions["C"].width = 32
    ws3.column_dimensions["I"].width = 32

    aging_fills = {
        label: PatternFill("solid", fgColor=bg) for label, (bg, _fg) in AGING_COLORS.items()
    }
    aging_fonts = {
        label: Font(bold=True, color=fg) for label, (_bg, fg) in AGING_COLORS.items()
    }

    # ---- Aging Report sheet ----
    ws4 = wb.create_sheet("Aging Report")
    aging_headers = ["Product", "Category", "Qty_ERP", "Aging"]
    for i, h in enumerate(aging_headers):
        ws4.cell(1, i + 1, h).font = bold
    for i, row in enumerate(aging_rows):
        for j, h in enumerate(aging_headers):
            cell = ws4.cell(i + 2, j + 1, row[h])
            if h == "Aging" and row[h] in aging_fills:
                cell.fill = aging_fills[row[h]]
                cell.font = aging_fonts[row[h]]
    for col, width in zip("ABCD", [42, 18, 14, 14]):
        ws4.column_dimensions[col].width = width

    # ---- Aging Report Detail sheet ----
    ws5 = wb.create_sheet("Aging Report Detail")
    detail_headers = ["Product", "Lot/Serial Number", "Category", "Lot Created Date", "Age (Months)", "Aging"]
    for i, h in enumerate(detail_headers):
        ws5.cell(1, i + 1, h).font = bold
    for i, row in enumerate(aging_detail_rows):
        for j, h in enumerate(detail_headers):
            value = row[h]
            cell = ws5.cell(i + 2, j + 1, value)
            if h == "Lot Created Date" and value is not None:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if h == "Aging" and value in aging_fills:
                cell.fill = aging_fills[value]
                cell.font = aging_fonts[value]
    for col, width in zip("ABCDEF", [42, 22, 18, 20, 14, 14]):
        ws5.column_dimensions[col].width = width

    wb.save(output_path)


def main():
    args = parse_args()
    date = infer_date(args.cleaned_csv, args.date)
    city = args.city
    opname_date = datetime.strptime(date, "%y%m%d")

    try:
        master = load_scope_and_prices(args.exclusion_xlsx)
    except odoo_client.OdooConfigError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
    whitelist = set(master.keys())
    product_lookup = build_product_lookup(whitelist)

    erp_lot_rows, erp_qty_by_product, abnormality_excluded_count, lot_tracked_products, clean_rows = load_cleaned_erp(args.cleaned_csv, city, product_lookup)

    wb_op = openpyxl.load_workbook(args.opname_xlsx, data_only=True)
    imei_sheet = f"Stock {city} {date} - IMEI"
    non_imei_sheet = f"Stock {city} {date} - Non IMEI"
    if imei_sheet not in wb_op.sheetnames or non_imei_sheet not in wb_op.sheetnames:
        print(f"ERROR: expected sheets {imei_sheet!r} / {non_imei_sheet!r} not found. "
              f"Available sheets: {wb_op.sheetnames}", file=sys.stderr)
        sys.exit(1)

    physical_lot_rows, duplicate_scans = load_physical_imei(wb_op, imei_sheet, product_lookup)
    physical_qty_nonimei = load_physical_non_imei(wb_op, non_imei_sheet, product_lookup)

    erp_lots = set(erp_lot_rows.keys())
    physical_lots = set(physical_lot_rows.keys())
    matched_lots = erp_lots & physical_lots
    erp_only_lots = erp_lots - physical_lots
    physical_only_lots = physical_lots - erp_lots
    product_mismatch_lots = {lot for lot in matched_lots if erp_lot_rows[lot]["Product"] != physical_lot_rows[lot]["Product"]}

    union_lots = erp_lots | physical_lots
    match_rate = (len(matched_lots) / len(union_lots)) if union_lots else None

    product_rows = build_product_summary(
        whitelist, master, erp_lot_rows, erp_qty_by_product,
        physical_lot_rows, physical_qty_nonimei, matched_lots, lot_tracked_products,
    )

    imei_tracked_rows = [p for p in product_rows if p["IMEI Match Rate (%)"] is not None]
    products_match = sum(1 for p in imei_tracked_rows if p["IMEI Match Rate (%)"] == 1)
    products_mismatch = sum(1 for p in imei_tracked_rows if p["IMEI Match Rate (%)"] < 1)

    summary_ctx = {
        "total_erp": len(erp_lots),
        "total_physical": len(physical_lots),
        "matched": len(matched_lots),
        "erp_only": len(erp_only_lots),
        "physical_only": len(physical_only_lots),
        "match_rate": match_rate,
        "accuracy_all": compute_accuracy(product_rows),
        "accuracy_ab": compute_accuracy(product_rows, groups={"A", "B"}),
        "products_match": products_match,
        "products_mismatch": products_mismatch,
        "discrepancy_all": sum(p["Discrepancy Value (IDR)"] for p in product_rows),
        "discrepancy_ab": sum(p["Discrepancy Value (IDR)"] for p in product_rows if p["New Group - July 25"] in {"A", "B"}),
    }

    mismatch_rows = []
    for lot in sorted(product_mismatch_lots):
        e, ph = erp_lot_rows[lot], physical_lot_rows[lot]
        mismatch_rows.append({
            "Lot/Serial Number": lot, "Match Status": "Product Mismatch",
            "Product_ERP": e["Product"], "Location": e["Location"], "Quantity": e["Quantity"],
            "Owner": e["Owner"], "Lot Created Date": e["Lot Created Date"], "Abnormality": e["Abnormality"],
            "Product_Physical": ph["Product"], "SKU": ph["SKU"], "Lokasi Gudang": ph["Lokasi Gudang"],
            "Box Ke-": ph["Box Ke-"], "Nomor Box": ph["Nomor Box"],
        })
    for lot in sorted(erp_only_lots):
        e = erp_lot_rows[lot]
        mismatch_rows.append({
            "Lot/Serial Number": lot, "Match Status": "Missing in Physical",
            "Product_ERP": e["Product"], "Location": e["Location"], "Quantity": e["Quantity"],
            "Owner": e["Owner"], "Lot Created Date": e["Lot Created Date"], "Abnormality": e["Abnormality"],
        })
    for lot in sorted(physical_only_lots):
        ph = physical_lot_rows[lot]
        mismatch_rows.append({
            "Lot/Serial Number": lot, "Match Status": "Missing in ERP",
            "Product_Physical": ph["Product"], "SKU": ph["SKU"], "Lokasi Gudang": ph["Lokasi Gudang"],
            "Box Ke-": ph["Box Ke-"], "Nomor Box": ph["Nomor Box"],
        })

    aging_rows, aging_skipped = build_aging_report(clean_rows, master, opname_date)
    aging_detail_rows = build_aging_detail(clean_rows, master, opname_date)

    output_path = args.output or str(Path(args.opname_xlsx).with_name(f"Stock Opname Report {city} {date}.xlsx"))
    write_report(output_path, city, date, summary_ctx, product_rows, mismatch_rows, aging_rows, aging_detail_rows)

    print(f"City: {city}  Date: {date}")
    print(f"ERP rows excluded (abnormality-flagged): {abnormality_excluded_count}")
    print(f"Total IMEIs in ERP: {summary_ctx['total_erp']}")
    print(f"Total IMEIs in Physical: {summary_ctx['total_physical']}")
    print(f"Matched: {summary_ctx['matched']}  ERP-only: {summary_ctx['erp_only']}  Physical-only: {summary_ctx['physical_only']}")
    print(f"Product Mismatch (same lot, different product name): {len(product_mismatch_lots)}")
    print(f"Physical-side duplicate scans: {len(duplicate_scans)}")
    print(f"IMEI Match Rate: {summary_ctx['match_rate']}")
    print(f"Stock Opname Accuracy (All): {summary_ctx['accuracy_all']}")
    print(f"Stock Opname Accuracy (A+B): {summary_ctx['accuracy_ab']}")
    print(f"IMEI-tracked products — Match: {products_match}  Mismatch: {products_mismatch}")
    print(f"Discrepancy Value (All): {summary_ctx['discrepancy_all']:,.0f} IDR")
    print(f"Discrepancy Value (A+B): {summary_ctx['discrepancy_ab']:,.0f} IDR")
    print(f"Aging report: {len(aging_rows)} product/bucket rows, {len(aging_detail_rows)} lot-level detail rows"
          + (f" ({aging_skipped} clean-row(s) skipped, no parseable Lot Created Date)" if aging_skipped else ""))
    print(f"Output written to: {output_path}")
    if duplicate_scans:
        print("NOTE: physical-side duplicate scans found (same lot scanned more than once) — data entry issue, not reflected above:")
        for lot, n in list(duplicate_scans.items())[:10]:
            print(f"   {lot}: scanned {n}x")


if __name__ == "__main__":
    main()
